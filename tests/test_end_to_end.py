from openpyxl import load_workbook

from vat_whitelist.cli import main
from vat_whitelist.inputs import read_rows


def test_reads_semicolon_csv_with_detected_columns(tmp_path):
    path = tmp_path / "k.csv"
    path.write_text("Firma;NIP;Rachunek\nA;526-025-02-74;\nB;;\n", encoding="utf-8")
    rows = read_rows(path)
    assert [(r.nip, r.account) for r in rows] == [("526-025-02-74", "")]


def test_reads_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Kontrahent", "NIP kontrahenta"])
    wb.active.append(["A", 5260250274])
    path = tmp_path / "k.xlsx"
    wb.save(path)
    assert read_rows(path)[0].nip == "5260250274"


def test_demo_run_writes_report(tmp_path, capsys):
    out = tmp_path / "raport.xlsx"
    code = main(["examples/kontrahenci.csv", "--demo", "-o", str(out), "--data", "2026-10-20"])
    assert code == 0
    wb = load_workbook(out)
    rows = list(wb["Wyniki"].iter_rows(values_only=True))
    assert rows[0][:3] == ("NIP (wejście)", "NIP", "Status VAT")
    statuses = [r[2] for r in rows[1:]]
    assert statuses.count("Czynny") == 5
    assert "Błędny NIP" in statuses
    assert "Raport:" in capsys.readouterr().out


def test_demo_run_csv_output(tmp_path):
    out = tmp_path / "raport.csv"
    assert main(["examples/kontrahenci.csv", "--demo", "-o", str(out)]) == 0
    assert out.read_text(encoding="utf-8-sig").startswith("NIP (wejście);NIP;Status VAT")


def test_accounts_go_to_separate_sheet(tmp_path):
    out = tmp_path / "raport.xlsx"
    main(["examples/kontrahenci.csv", "--demo", "-o", str(out)])
    wb = load_workbook(out)
    header = next(wb["Wyniki"].iter_rows(values_only=True))
    first = list(wb["Wyniki"].iter_rows(min_row=2, values_only=True))[0]
    assert first[header.index("Rachunków w wykazie")] == 1
    accounts = list(wb["Rachunki"].iter_rows(min_row=2, values_only=True))
    # duplicate NIP in the input file is listed once
    assert [a[0] for a in accounts].count("0004182736") == 1
    assert (
        "0004182736",
        "HURTOWNIA ALFA SP. Z O.O. (DEMO)",
        "60102010260000042270201111",
    ) in accounts
