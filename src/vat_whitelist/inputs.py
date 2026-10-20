"""Reading the list of contractors from CSV or Excel."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from .checker import InputRow

_NIP_HEADERS = ("nip", "nip kontrahenta", "vat id", "tax id")
_ACCOUNT_HEADERS = ("rachunek", "konto", "nr rachunku", "numer rachunku", "iban", "account")


class InputError(ValueError):
    pass


def read_rows(
    path: str | Path, nip_column: str | None = None, account_column: str | None = None
) -> list[InputRow]:
    path = Path(path)
    if not path.exists():
        raise InputError(f"Nie ma pliku {path}")
    table = _read_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") else _read_csv(path)
    if not table:
        raise InputError(f"Plik {path} jest pusty")

    header = [str(h or "").strip() for h in table[0]]
    nip_idx = _find_column(header, nip_column, _NIP_HEADERS)
    if nip_idx is None:
        raise InputError(
            f"Nie znalazłem kolumny z NIP-em w nagłówku {header}. Podaj ją przez --kolumna-nip."
        )
    acc_idx = _find_column(header, account_column, _ACCOUNT_HEADERS)
    if account_column and acc_idx is None:
        raise InputError(f"Nie ma kolumny {account_column!r} w nagłówku {header}")

    rows = []
    for record in table[1:]:
        nip = _cell(record, nip_idx)
        if not nip:
            continue
        rows.append(
            InputRow(nip=nip, account=_cell(record, acc_idx) if acc_idx is not None else "")
        )
    return rows


def _find_column(header: list[str], wanted: str | None, candidates: tuple[str, ...]) -> int | None:
    lowered = [h.lower() for h in header]
    if wanted:
        return lowered.index(wanted.lower()) if wanted.lower() in lowered else None
    for i, name in enumerate(lowered):
        if name in candidates:
            return i
    return None


def _cell(record: list, idx: int) -> str:
    if idx >= len(record) or record[idx] is None:
        return ""
    value = record[idx]
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_csv(path: Path) -> list[list[str]]:
    text = path.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect) if any(row)]


def _read_xlsx(path: Path) -> list[list]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    return [list(row) for row in sheet.iter_rows(values_only=True) if any(row)]
