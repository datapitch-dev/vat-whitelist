"""Excel and CSV reports."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .checker import STATUS_PENDING, Result, RunStats

COLUMNS = [
    ("NIP (wejście)", lambda r: r.input_nip, 16),
    ("NIP", lambda r: r.nip, 13),
    ("Status VAT", lambda r: r.status, 18),
    ("Nazwa", lambda r: r.name, 40),
    ("REGON", lambda r: r.regon, 12),
    ("KRS", lambda r: r.krs, 12),
    ("Adres", lambda r: r.address, 40),
    ("Data rejestracji", lambda r: r.registration_date, 14),
    ("Data wykreślenia", lambda r: r.removal_date, 14),
    ("Rachunek z pliku", lambda r: r.input_account, 30),
    ("Rachunek na liście", lambda r: r.account_on_list, 16),
    ("Rachunków w wykazie", lambda r: len(r.accounts) if r.accounts else None, 12),
    ("Rachunki wirtualne", lambda r: "tak" if r.has_virtual_accounts else "", 10),
    ("Id zapytania (dowód)", lambda r: r.request_id, 22),
    ("Uwagi", lambda r: r.note, 40),
]

_GREEN = PatternFill("solid", fgColor="C6EFCE")
_YELLOW = PatternFill("solid", fgColor="FFEB9C")
_RED = PatternFill("solid", fgColor="FFC7CE")
_GREY = PatternFill("solid", fgColor="E7E6E6")


def _fill(result: Result) -> PatternFill:
    if result.status == STATUS_PENDING:
        return _GREY
    if result.account_on_list in ("NIE", "Błędny numer"):
        return _RED
    if result.is_active:
        return _YELLOW if result.account_on_list == "Sprawdź ręcznie" else _GREEN
    return _YELLOW if result.status == "Zwolniony" else _RED


def summary(results: list[Result]) -> Counter:
    counts = Counter(r.status for r in results)
    for r in results:
        if r.account_on_list:
            counts[f"Rachunek: {r.account_on_list}"] += 1
    return counts


def write_xlsx(results: list[Result], stats: RunStats, path: str | Path, date: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wyniki"
    ws.append([name for name, _, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for result in results:
        ws.append([get(result) for _, get, _ in COLUMNS])
        fill = _fill(result)
        for cell in ws[ws.max_row][:3]:
            cell.fill = fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    # Large companies list hundreds of accounts, so they get their own sheet, one per row.
    accounts = wb.create_sheet("Rachunki")
    accounts.append(["NIP", "Nazwa", "Rachunek z wykazu"])
    for cell in accounts[1]:
        cell.font = Font(bold=True)
    for width, col in ((13, "A"), (40, "B"), (30, "C")):
        accounts.column_dimensions[col].width = width
    seen = set()
    for result in results:
        if result.nip in seen:
            continue
        seen.add(result.nip)
        for account in result.accounts:
            accounts.append([result.nip, result.name, account])
    accounts.freeze_panes = "A2"
    accounts.auto_filter.ref = accounts.dimensions

    info = wb.create_sheet("Podsumowanie")
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 60
    info.append(["Stan na dzień", date])
    info.append(["Kontrahentów w pliku", len(results)])
    info.append(["Zapytań do API w tym uruchomieniu", stats.requests_used])
    info.append(["Odczytanych z pamięci podręcznej", stats.from_cache])
    if stats.stopped_reason:
        info.append(["Uwaga", stats.stopped_reason])
    info.append([])
    info.append(["Wynik", "Liczba"])
    info[info.max_row][0].font = info[info.max_row][1].font = Font(bold=True)
    for key, count in sorted(summary(results).items()):
        info.append([key, count])
    info.append([])
    info.append(["Źródło", "Wykaz podatników VAT, Ministerstwo Finansów (wl-api.mf.gov.pl)"])
    wb.save(path)


def write_csv(results: list[Result], path: str | Path) -> None:
    # Semicolon + BOM so Polish Excel opens it correctly with a double click.
    with Path(path).open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow([name for name, _, _ in COLUMNS])
        for result in results:
            writer.writerow(["" if (v := get(result)) is None else v for _, get, _ in COLUMNS])
